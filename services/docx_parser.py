from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import NamedTuple

from docx import Document
from docx.oxml.ns import qn

from core.logging import get_logger

logger = get_logger(__name__)


class TextBlock(NamedTuple):
    text: str
    style: str  # paragraph style name, e.g. "Heading 2", "Normal"


@dataclass
class ParsedDocument:
    blocks: list[TextBlock] = field(default_factory=list)
    # Flat joined text from tables (labels + values)
    table_lines: list[str] = field(default_factory=list)
    # Raw date hint extracted by parsers (e.g. "2 дні тому", "16 квітня 2026 року")
    raw_date_hint: str = ""

    @property
    def all_lines(self) -> list[str]:
        """All non-empty text lines from paragraphs + tables."""
        lines: list[str] = []
        for block in self.blocks:
            for line in block.text.splitlines():
                stripped = line.strip()
                if stripped:
                    lines.append(stripped)
        lines.extend(self.table_lines)
        return lines


def _extract_table_lines(doc: Document) -> list[str]:
    """Return key-value style lines from all tables."""
    lines: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells]
            # Deduplicate merged cells that appear twice
            seen: set[str] = set()
            unique: list[str] = []
            for t in cell_texts:
                if t and t not in seen:
                    seen.add(t)
                    unique.append(t)
            if unique:
                lines.append("\t".join(unique))
    return lines


def parse_docx(path: Path) -> ParsedDocument:
    """
    Extract all text blocks and table content from a .docx file.
    Returns a ParsedDocument. Never raises — logs errors instead.
    """
    try:
        doc = Document(str(path))
    except Exception as exc:
        logger.error("Failed to open %s: %s", path.name, exc)
        return ParsedDocument()

    blocks: list[TextBlock] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(TextBlock(text=text, style=para.style.name))

    try:
        table_lines = _extract_table_lines(doc)
    except Exception as exc:
        logger.warning("Table extraction failed for %s: %s", path.name, exc)
        table_lines = []

    return ParsedDocument(blocks=blocks, table_lines=table_lines)
