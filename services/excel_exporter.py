from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.logging import get_logger
from services.resume_extractor import ResumeData

logger = get_logger(__name__)

COLUMNS: list[str] = [
    "Дата публікації",
    "ПІБ",
    "Номер телефону",
    "Вік",
    "Посада",
    "Досвід роботи",
    "Джерело",
]

# Header styling
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="2F5496")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Minimum column widths (characters)
_MIN_COL_WIDTHS: dict[str, int] = {
    "Дата публікації": 18,
    "Посада": 45,
    "ПІБ": 28,
    "Вік": 8,
    "Джерело": 16,
    "Номер телефону": 22,
    "Досвід роботи": 45,
}

def _auto_fit_column(ws, col_idx: int, header: str, values: list[str]) -> None:
    letter = get_column_letter(col_idx)
    max_len = max(
        [len(header)] + [len(str(v)) for v in values]
    )
    min_width = _MIN_COL_WIDTHS.get(header, 15)
    ws.column_dimensions[letter].width = max(min_width, min(max_len + 4, 60))


def build_excel(resumes: list[ResumeData], output_path: Path) -> None:
    """Write a list of ResumeData objects to an .xlsx file.

    Rows are written in the same order as the input list, which is the order
    in which the user uploaded the resumes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.freeze_panes = "A2"  # freeze header row

    # Write header
    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    ws.row_dimensions[1].height = 22

    # Write data rows
    row_values: dict[str, list[str]] = {col: [] for col in COLUMNS}
    for row_idx, resume in enumerate(resumes, start=2):
        date_str = (
            resume.publication_date.strftime("%d.%m.%Y")
            if resume.publication_date
            else ""
        )
        values = [
            date_str,
            resume.name,
            resume.phone,
            resume.age,
            resume.positions,
            resume.work_experience,
            resume.source,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            row_values[COLUMNS[col_idx - 1]].append(str(value))

    # Auto-fit columns
    for col_idx, header in enumerate(COLUMNS, start=1):
        _auto_fit_column(ws, col_idx, header, row_values[header])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    logger.info("Excel saved: %s (%d rows)", output_path, len(resumes))
